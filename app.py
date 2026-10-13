import calendar
import io
import json
import math
import os
import re
import secrets
import tempfile
import time
from datetime import date, datetime, timedelta, timezone
from functools import wraps

import click
from flask import (Flask, Response, abort, flash, g, redirect, render_template,
                   request, session, url_for)
from werkzeug.security import check_password_hash, generate_password_hash

import db

# ------------------------------------------------------------------ settings
FEE_PER_YEAR = 250      # Taka per year = 24 issues  ->  250/24 Taka per issue
ISSUES_PER_YEAR = 24    # fortnightly magazine
FY_START_MONTH = 7      # financial year July - June
PER_PAGE = 50
ROUND_TO = 1            # "amount to pay" is rounded UP to a multiple of this (1 = whole Taka, 5 = nearest 5 Taka)
# Can a local Jama'at user record new payments? (False = only the central admin records money)
JAMAAT_USERS_CAN_ADD_PAYMENTS = False

app = Flask(__name__)
app.secret_key = os.environ.get("SECRET_KEY", "change-this-secret-key-before-going-live")
if os.environ.get("RENDER"):                       # running online on Render (https)
    app.config.update(SESSION_COOKIE_SECURE=True, SESSION_COOKIE_HTTPONLY=True, SESSION_COOKIE_SAMESITE="Lax")
app.config["MAX_CONTENT_LENGTH"] = 5 * 1024 * 1024      # biggest Excel upload: 5 MB
app.teardown_appcontext(db.close_db)


# ------------------------------------------------------------------ date helpers
def today():
    # Bangladesh time (UTC+6) even when the server (e.g. Render) runs on UTC
    return datetime.now(timezone(timedelta(hours=6))).date()


def fy_of(d):
    """Financial-year start year of a date (2025-07-01 .. 2026-06-30 -> 2025)."""
    if isinstance(d, str):
        d = datetime.strptime(d[:10], "%Y-%m-%d").date()
    return d.year if d.month >= FY_START_MONTH else d.year - 1


def fy_label(y):
    return f"{y}-{str(y + 1)[2:]}"


def fy_range(y):
    return f"{y}-07-01", f"{y + 1}-06-30"


def add_months(d, n):
    m = d.month - 1 + n
    y = d.year + m // 12
    m = m % 12 + 1
    return date(y, m, min(d.day, calendar.monthrange(y, m)[1]))


def parse_date(s):
    try:
        return datetime.strptime(s, "%Y-%m-%d").date()
    except (TypeError, ValueError):
        return None


@app.template_filter("dmy")
def dmy(s):
    d = parse_date((s or "")[:10]) if not isinstance(s, date) else s
    return d.strftime("%d-%m-%Y") if d else ""


@app.template_filter("money")
def money(v):
    v = float(v or 0)
    return f"{v:,.0f}" if v == int(v) else f"{v:,.2f}"


# ------------------------------------------------------------------ amount to pay
def renewal_quote(expiry, extra_years=0):
    """What a subscriber must pay so the subscription ends exactly on a 30 June.

    * The price is FEE_PER_YEAR / 24 per issue (24 issues = 1 year = 250 Taka).
    * Cover is needed from the first day not yet covered: the day after the current expiry
      (if still running) or today (if already expired).
    * It is topped up to the 30 June that ends the financial year (July-June) containing that day,
      so from the next July everybody simply pays the round yearly fee.
    * extra_years adds that many more full years (24 issues each) after that 30 June.
    Returns dict(amount, issues, first, until)."""
    t = today()
    exp = parse_date((expiry or "")[:10])
    first = exp + timedelta(days=1) if exp and exp >= t else t
    fy = fy_of(first)
    fy_start, fy_end = date(fy, 7, 1), date(fy + 1, 6, 30)
    days = (fy_end - first).days + 1                  # days still to be covered in that year
    fy_days = (fy_end - fy_start).days + 1            # 365 or 366
    issues = math.ceil(round(days * ISSUES_PER_YEAR / fy_days, 6)) + ISSUES_PER_YEAR * extra_years
    amount = math.ceil(round(issues * FEE_PER_YEAR / ISSUES_PER_YEAR / ROUND_TO, 6)) * ROUND_TO
    return {"amount": amount, "issues": issues, "first": first,
            "until": date(fy + 1 + extra_years, 6, 30)}


# ------------------------------------------------------------------ users / permissions
@app.before_request
def load_user():
    g.user = None
    uid = session.get("uid")
    if uid:
        g.user = db.get_db().execute(
            """SELECT u.id, u.username, u.role, u.jamaat_id, j.name AS jamaat
               FROM users u LEFT JOIN jamaats j ON j.id = u.jamaat_id WHERE u.id=?""",
            (uid,)).fetchone()
        if g.user is None:
            session.clear()


def is_admin():
    return bool(g.get("user") and g.user["role"] == "admin")


def can_pay():
    return is_admin() or (g.get("user") is not None and JAMAAT_USERS_CAN_ADD_PAYMENTS)


def login_required(f):
    @wraps(f)
    def w(*a, **k):
        if g.user is None:
            return redirect(url_for("login", next=request.path))
        return f(*a, **k)
    return w


def admin_required(f):
    @wraps(f)
    def w(*a, **k):
        if g.user is None:
            return redirect(url_for("login", next=request.path))
        if not is_admin():
            abort(403)
        return f(*a, **k)
    return w


def scope(where, args, col="jamaat_id"):
    """Local Jama'at users only ever see their own Jama'at."""
    if not is_admin():
        where.append(f"{col}=?")
        args.append(g.user["jamaat_id"])


def load_sub(sid):
    s = db.get_db().execute("SELECT * FROM v_subscribers WHERE id=?", (sid,)).fetchone()
    if not s or (not is_admin() and s["jamaat_id"] != g.user["jamaat_id"]):
        abort(404)
    return s


# ------------------------------------------------------------------ sorting
SORTS = {
    "id": [("id", "dir")],
    "local": [("(local_id IS NULL OR local_id='')", "fix"), ("CAST(local_id AS INTEGER)", "dir"),
              ("local_id", "dir"), ("lower(name)", "fix")],
    "name": [("lower(name)", "dir"), ("id", "fix")],
    "jamaat": [("(jamaat IS NULL)", "fix"), ("lower(jamaat)", "dir"),
               ("lower(name)", "fix")],
    "expiry": [("(expiry IS NULL)", "fix"), ("expiry", "dir"), ("lower(name)", "fix")],
}
SORT_LABELS = [("id", "ID"), ("local", "Local ID"), ("name", "Name"),
               ("jamaat", "Jama'at"), ("expiry", "Expiry date")]


def get_sort(default="id"):
    g.default_sort = default
    key = request.args.get("sort", default)
    if key not in SORTS:
        key = default
    return key, ("desc" if request.args.get("dir") == "desc" else "asc")


def order_sql(key, d, lead=None):
    parts = [lead] if lead else []
    for expr, mode in SORTS[key]:
        parts.append(expr + (" DESC" if (mode == "dir" and d == "desc") else " ASC"))
    return " ORDER BY " + ", ".join(parts)


def sort_url(key):
    args = request.args.to_dict()
    cur = args.get("sort", g.get("default_sort"))
    d = args.get("dir", "asc")
    args["sort"] = key
    args["dir"] = "desc" if (cur == key and d == "asc") else "asc"
    args.pop("page", None)
    return url_for(request.endpoint, **args)


def sort_mark(key):
    cur = request.args.get("sort", g.get("default_sort"))
    if cur == key:
        return " ▲" if request.args.get("dir", "asc") == "asc" else " ▼"
    return ""


def page_url(n):
    args = request.args.to_dict()
    args["page"] = n
    return url_for(request.endpoint, **args)


@app.context_processor
def inject():
    return dict(fy_label=fy_label, FEE=FEE_PER_YEAR, today=today().isoformat(),
                current_fy=fy_of(today()), user=g.get("user"), is_admin=is_admin(),
                can_pay=can_pay(), sort_url=sort_url, sort_mark=sort_mark, page_url=page_url,
                sort_labels=SORT_LABELS)


# ------------------------------------------------------------------ misc helpers
def xlsx_response(title, headers, rows, filename):
    from openpyxl import Workbook
    from openpyxl.styles import Font
    wb = Workbook()
    ws = wb.active
    ws.title = title[:30]
    ws.append(headers)
    for c in ws[1]:
        c.font = Font(bold=True)
    for r in rows:
        ws.append(list(r))
    for col in ws.columns:
        ws.column_dimensions[col[0].column_letter].width = min(
            45, max(10, max(len(str(c.value or "")) for c in col) + 2))
    buf = io.BytesIO()
    wb.save(buf)
    return Response(buf.getvalue(), headers={
        "Content-Disposition": f"attachment; filename={filename}.xlsx",
        "Content-Type": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"})


def jamaat_list():
    if not is_admin():
        return db.get_db().execute("SELECT id, name FROM jamaats WHERE id=?",
                                   (g.user["jamaat_id"],)).fetchall()
    return db.get_db().execute("SELECT id, name FROM jamaats ORDER BY name").fetchall()


def get_or_create_jamaat(name):
    name = (name or "").strip()
    if not name:
        return None
    con = db.get_db()
    r = con.execute("SELECT id FROM jamaats WHERE lower(name)=lower(?)", (name,)).fetchone()
    if r:
        return r["id"]
    return con.execute("INSERT INTO jamaats(name) VALUES (?)", (name,)).lastrowid


def jam_filter_value():
    """Jamaat filter chosen in the URL (always forced to own Jama'at for local users)."""
    if not is_admin():
        return str(g.user["jamaat_id"])
    return request.args.get("jamaat", "")


# ------------------------------------------------------------------ login / account
@app.route("/login", methods=["GET", "POST"])
def login():
    if request.method == "POST":
        u = db.get_db().execute("SELECT * FROM users WHERE lower(username)=lower(?)",
                                (request.form["username"].strip(),)).fetchone()
        if u and check_password_hash(u["password_hash"], request.form["password"]):
            session.clear()
            session["uid"] = u["id"]
            nxt = request.args.get("next") or url_for("dashboard")
            return redirect(nxt if nxt.startswith("/") else url_for("dashboard"))
        flash("Wrong username or password.", "danger")
    return render_template("login.html")


@app.route("/logout")
def logout():
    session.clear()
    return redirect(url_for("login"))


@app.route("/account/password", methods=["GET", "POST"])
@login_required
def account_password():
    if request.method == "POST":
        con = db.get_db()
        u = con.execute("SELECT * FROM users WHERE id=?", (g.user["id"],)).fetchone()
        new = request.form.get("new_password", "")
        if not check_password_hash(u["password_hash"], request.form.get("old_password", "")):
            flash("Current password is wrong.", "danger")
        elif len(new) < 6:
            flash("New password must be at least 6 characters.", "danger")
        elif new != request.form.get("new_password2"):
            flash("The two new passwords do not match.", "danger")
        else:
            con.execute("UPDATE users SET password_hash=? WHERE id=?",
                        (generate_password_hash(new), u["id"]))
            con.commit()
            flash("Password changed.", "success")
            return redirect(url_for("dashboard"))
    return render_template("account.html")


# ------------------------------------------------------------------ user management (admin)
@app.route("/users", methods=["GET", "POST"])
@admin_required
def users_page():
    con = db.get_db()
    if request.method == "POST":
        username = request.form.get("username", "").strip()
        pw = request.form.get("password", "")
        role = request.form.get("role", "")
        jam = request.form.get("jamaat_id", "")
        err = None
        if not username:
            err = "Username is required."
        elif len(pw) < 6:
            err = "Password must be at least 6 characters."
        elif role not in ("admin", "jamaat"):
            err = "Choose a role."
        elif role == "jamaat" and not (jam.isdigit() and con.execute(
                "SELECT 1 FROM jamaats WHERE id=?", (int(jam),)).fetchone()):
            err = "Choose the Jama'at for this local user."
        elif con.execute("SELECT 1 FROM users WHERE lower(username)=lower(?)", (username,)).fetchone():
            err = "That username already exists."
        if err:
            flash(err, "danger")
        else:
            con.execute("INSERT INTO users(username,password_hash,role,jamaat_id) VALUES (?,?,?,?)",
                        (username, generate_password_hash(pw), role, int(jam) if role == "jamaat" else None))
            con.commit()
            flash(f"User '{username}' created.", "success")
        return redirect(url_for("users_page"))
    users = con.execute(
        """SELECT u.id, u.username, u.role, j.name AS jamaat FROM users u
           LEFT JOIN jamaats j ON j.id=u.jamaat_id ORDER BY u.role, j.name, u.username""").fetchall()
    return render_template("users.html", users=users,
                           jamaats=con.execute("SELECT id,name FROM jamaats ORDER BY name").fetchall())


@app.route("/users/<int:uid>/password", methods=["POST"])
@admin_required
def user_password(uid):
    pw = request.form.get("new_password", "")
    if len(pw) < 6:
        flash("Password must be at least 6 characters.", "danger")
    else:
        con = db.get_db()
        con.execute("UPDATE users SET password_hash=? WHERE id=?", (generate_password_hash(pw), uid))
        con.commit()
        flash("Password reset.", "success")
    return redirect(url_for("users_page"))


@app.route("/users/<int:uid>/delete", methods=["POST"])
@admin_required
def user_delete(uid):
    con = db.get_db()
    u = con.execute("SELECT * FROM users WHERE id=?", (uid,)).fetchone()
    if not u:
        abort(404)
    if uid == g.user["id"]:
        flash("You cannot delete your own login.", "danger")
    elif u["role"] == "admin" and con.execute(
            "SELECT COUNT(*) FROM users WHERE role='admin'").fetchone()[0] <= 1:
        flash("You cannot delete the last admin.", "danger")
    else:
        con.execute("DELETE FROM users WHERE id=?", (uid,))
        con.commit()
        flash(f"User '{u['username']}' deleted.", "warning")
    return redirect(url_for("users_page"))


# ------------------------------------------------------------------ dashboard
@app.route("/")
@login_required
def dashboard():
    con = db.get_db()
    w, a = [], []
    scope(w, a)
    ws = (" WHERE " + " AND ".join(w)) if w else ""
    counts = {r["status"]: r["n"] for r in con.execute(
        f"SELECT status, COUNT(*) n FROM v_subscribers{ws} GROUP BY status", a)}
    t = today()
    in30 = con.execute(
        f"SELECT COUNT(*) FROM v_subscribers WHERE status='Valid' AND expiry <= ?"
        + (" AND " + " AND ".join(w) if w else ""),
        [add_months(t, 1).isoformat()] + a).fetchone()[0]
    y = fy_of(t)
    fa, fb = fy_range(y)
    sc, sa = "", []
    if not is_admin():
        sc, sa = " AND s.jamaat_id=?", [g.user["jamaat_id"]]
    fy_total = con.execute(
        f"""SELECT COALESCE(SUM(p.amount),0) s, COUNT(*) n FROM payments p
            JOIN subscribers s ON s.id=p.subscriber_id WHERE p.paid_date BETWEEN ? AND ?{sc}""",
        [fa, fb] + sa).fetchone()
    recent = con.execute(
        f"""SELECT p.*, s.name, s.local_id, j.name AS jamaat FROM payments p
            JOIN subscribers s ON s.id=p.subscriber_id LEFT JOIN jamaats j ON j.id=s.jamaat_id
            WHERE 1=1{sc} ORDER BY p.paid_date DESC, p.id DESC LIMIT 10""", sa).fetchall()
    return render_template("dashboard.html", counts=counts, in30=in30, fy=y, fy_total=fy_total,
                           recent=recent)


# ------------------------------------------------------------------ subscribers
@app.route("/subscribers")
@login_required
def subscribers():
    q = request.args.get("q", "").strip()
    jam = jam_filter_value()
    status = request.args.get("status", "")
    sort, d = get_sort("id")
    try:
        page = max(1, int(request.args.get("page", 1)))
    except ValueError:
        page = 1
    where, args = [], []
    scope(where, args)
    if q:
        where.append("(name LIKE ? OR father_name LIKE ? OR mobile LIKE ? OR address LIKE ? "
                     "OR local_id LIKE ? OR CAST(id AS TEXT)=?)")
        like = f"%{q}%"
        args += [like, like, like, like, like, q]
    if jam == "none":
        where.append("jamaat_id IS NULL")
    elif jam and is_admin():
        where.append("jamaat_id=?")
        args.append(jam)
    if status:
        where.append("status=?")
        args.append(status)
    w = " WHERE " + " AND ".join(where) if where else ""
    con = db.get_db()
    total = con.execute(f"SELECT COUNT(*) FROM v_subscribers{w}", args).fetchone()[0]
    rows = con.execute(f"SELECT * FROM v_subscribers{w}{order_sql(sort, d)} LIMIT ? OFFSET ?",
                       args + [PER_PAGE, (page - 1) * PER_PAGE]).fetchall()
    pages = max(1, -(-total // PER_PAGE))
    return render_template("subscribers.html", rows=rows, total=total, page=page, pages=pages,
                           q=q, jam=jam, status=status, sort=sort, dir=d, jamaats=jamaat_list())


def _sub_form_data(existing=None):
    f = request.form
    d = dict(
        name=f.get("name", "").strip(), local_id=f.get("local_id", "").strip() or None,
        father_name=f.get("father_name", "").strip() or None,
        address=f.get("address", "").strip() or None, mobile=f.get("mobile", "").strip() or None,
        email=f.get("email", "").strip() or None, by_hand=1 if f.get("by_hand") else 0,
        active=1 if f.get("active") else 0, notes=f.get("notes", "").strip() or None)
    if is_admin():
        d["jamaat"] = f.get("jamaat_name", "").strip()
        d["jamaat_id"] = get_or_create_jamaat(d["jamaat"])
    else:                                   # local users can never move a member to another Jama'at
        d["jamaat_id"] = g.user["jamaat_id"]
        d["jamaat"] = g.user["jamaat"]
    return d


def local_id_conflict(jamaat_id, local_id, exclude=None):
    if not local_id or jamaat_id is None:
        return None
    return db.get_db().execute(
        "SELECT id, name FROM subscribers WHERE jamaat_id=? AND lower(local_id)=lower(?) AND id IS NOT ?",
        (jamaat_id, local_id, exclude)).fetchone()


@app.route("/subscribers/new", methods=["GET", "POST"])
@login_required
def subscriber_new():
    if request.method == "POST":
        d = _sub_form_data()
        err = None
        if not d["name"]:
            err = "Name is required."
        else:
            c = local_id_conflict(d["jamaat_id"], d["local_id"])
            if c:
                err = f"Local ID {d['local_id']} is already used in this Jama'at by #{c['id']} {c['name']}."
        if err:
            flash(err, "danger")
            return render_template("subscriber_form.html", s=d, jamaats=jamaat_list(), new=True)
        con = db.get_db()
        cur = con.execute(
            """INSERT INTO subscribers(local_id,name,father_name,jamaat_id,address,mobile,email,by_hand,active,notes)
               VALUES (:local_id,:name,:father_name,:jamaat_id,:address,:mobile,:email,:by_hand,:active,:notes)""", d)
        con.commit()
        if can_pay():
            flash(f"Subscriber added with ID {cur.lastrowid}. Now record the first payment.", "success")
            return redirect(url_for("payment_new", sid=cur.lastrowid))
        flash(f"Subscriber added with ID {cur.lastrowid}.", "success")
        return redirect(url_for("subscriber_view", sid=cur.lastrowid))
    return render_template("subscriber_form.html", jamaats=jamaat_list(), new=True,
                           s={"active": 1, "by_hand": 0, "jamaat": None if is_admin() else g.user["jamaat"]})


@app.route("/subscribers/<int:sid>/edit", methods=["GET", "POST"])
@login_required
def subscriber_edit(sid):
    s = load_sub(sid)
    con = db.get_db()
    if request.method == "POST":
        d = _sub_form_data()
        if not is_admin():
            d["jamaat_id"] = s["jamaat_id"]
        err = None
        if not d["name"]:
            err = "Name is required."
        else:
            c = local_id_conflict(d["jamaat_id"], d["local_id"], exclude=sid)
            if c:
                err = f"Local ID {d['local_id']} is already used in this Jama'at by #{c['id']} {c['name']}."
        if err:
            flash(err, "danger")
            d["id"] = sid
            return render_template("subscriber_form.html", s=d, jamaats=jamaat_list(), new=False)
        d["id"] = sid
        con.execute("""UPDATE subscribers SET local_id=:local_id, name=:name, father_name=:father_name,
                       jamaat_id=:jamaat_id, address=:address, mobile=:mobile, email=:email,
                       by_hand=:by_hand, active=:active, notes=:notes WHERE id=:id""", d)
        con.commit()
        flash("Subscriber updated.", "success")
        return redirect(url_for("subscriber_view", sid=sid))
    return render_template("subscriber_form.html", s=s, jamaats=jamaat_list(), new=False)


@app.route("/subscribers/<int:sid>")
@login_required
def subscriber_view(sid):
    s = load_sub(sid)
    pays = db.get_db().execute("SELECT * FROM payments WHERE subscriber_id=? ORDER BY paid_date, id",
                               (sid,)).fetchall()
    by_fy = {}
    for p in pays:
        by_fy.setdefault(fy_of(p["paid_date"]), []).append(p)
    return render_template("subscriber_view.html", s=s, by_fy=dict(sorted(by_fy.items(), reverse=True)))


# ------------------------------------------------------------------ payments
def _current_expiry(sid, exclude_payment=None):
    r = db.get_db().execute(
        "SELECT MAX(expiry_date) FROM payments WHERE subscriber_id=? AND id IS NOT ?",
        (sid, exclude_payment)).fetchone()
    return r[0]


def _save_payment(sid, pid=None):
    f = request.form
    paid = parse_date(f.get("paid_date"))
    exp = parse_date(f.get("expiry_date"))
    try:
        amount = float(f.get("amount", "").replace(",", ""))
    except ValueError:
        amount = None
    if not paid or not exp or not amount or amount <= 0:
        flash("Payment date, a positive amount and an expiry date are required.", "danger")
        return False
    vals = (f.get("receipt_no", "").strip() or None, f.get("book_no", "").strip() or None,
            paid.isoformat(), amount, exp.isoformat(), f.get("note", "").strip() or None)
    con = db.get_db()
    if vals[0]:                                   # same ID + Receipt No + Book No = double entry
        dup = con.execute("""SELECT id FROM payments WHERE subscriber_id=? AND lower(receipt_no)=lower(?)
                              AND lower(COALESCE(book_no,''))=lower(?) AND id IS NOT ?""",
                          (sid, vals[0], vals[1] or "", pid)).fetchone()
        if dup:
            flash("Blocked: this subscriber already has a payment with the same Receipt No and Book No.", "danger")
            return False
    if pid:
        con.execute("""UPDATE payments SET receipt_no=?, book_no=?, paid_date=?, amount=?,
                       expiry_date=?, note=? WHERE id=?""", vals + (pid,))
    else:
        con.execute("""INSERT INTO payments(receipt_no,book_no,paid_date,amount,expiry_date,note,subscriber_id)
                       VALUES (?,?,?,?,?,?,?)""", vals + (sid,))
    con.commit()
    return True


@app.route("/subscribers/<int:sid>/payments/new", methods=["GET", "POST"])
@login_required
def payment_new(sid):
    if not can_pay():
        abort(403)
    s = load_sub(sid)
    if request.method == "POST" and _save_payment(sid):
        flash("Payment recorded.", "success")
        return redirect(url_for("subscriber_view", sid=sid))
    cur = _current_expiry(sid)
    quotes = []
    for extra in (0, 1):                      # top-up to the 30 June, and the same plus one more full year
        q = renewal_quote(cur, extra)
        quotes.append({"amount": q["amount"], "issues": q["issues"], "until": q["until"].isoformat(),
                       "until_dmy": q["until"].strftime("%d-%m-%Y"), "extra": extra})
    return render_template("payment_form.html", s=s, p=request.form or {"paid_date": today().isoformat(),
                           "amount": FEE_PER_YEAR}, current_expiry=cur, edit=False, quotes=quotes)


@app.route("/payments/<int:pid>/edit", methods=["GET", "POST"])
@admin_required
def payment_edit(pid):
    con = db.get_db()
    p = con.execute("SELECT * FROM payments WHERE id=?", (pid,)).fetchone()
    if not p:
        abort(404)
    sid = p["subscriber_id"]
    if request.method == "POST" and _save_payment(sid, pid):
        flash("Payment updated.", "success")
        return redirect(url_for("subscriber_view", sid=sid))
    s = con.execute("SELECT * FROM v_subscribers WHERE id=?", (sid,)).fetchone()
    return render_template("payment_form.html", s=s, p=request.form or p, quotes=[],
                           current_expiry=_current_expiry(sid, pid), edit=True, pid=pid)


@app.route("/payments/<int:pid>/delete", methods=["POST"])
@admin_required
def payment_delete(pid):
    con = db.get_db()
    p = con.execute("SELECT subscriber_id FROM payments WHERE id=?", (pid,)).fetchone()
    if not p:
        abort(404)
    con.execute("DELETE FROM payments WHERE id=?", (pid,))
    con.commit()
    flash("Payment deleted.", "warning")
    return redirect(url_for("subscriber_view", sid=p["subscriber_id"]))


# ------------------------------------------------------------------ reports
@app.route("/reports/jamaat")
@login_required
def report_jamaat():
    """Jamaat-wise subscriber list with expiry dates."""
    jam = jam_filter_value()
    status = request.args.get("status", "")
    sort, d = get_sort("name")
    where, args = [], []
    scope(where, args)
    if jam == "none":
        where.append("jamaat_id IS NULL")
    elif jam and is_admin():
        where.append("jamaat_id=?")
        args.append(jam)
    if status == "current":
        where.append("status='Valid'")
    elif status:
        where.append("status=?")
        args.append(status)
    w = " WHERE " + " AND ".join(where) if where else ""
    lead = "(jamaat IS NULL) ASC, lower(jamaat) " + ("DESC" if (sort == "jamaat" and d == "desc") else "ASC")
    rows = db.get_db().execute(f"SELECT * FROM v_subscribers{w}{order_sql(sort, d, lead)}", args).fetchall()
    if request.args.get("format") == "xlsx":
        return xlsx_response("Jamaat-wise", ["Jamaat", "ID", "Local ID", "Name", "Father's name", "Mobile",
                             "Address", "Delivery", "Status", "Expiry date", "Total paid"],
                             [(r["jamaat"] or "(none)", r["id"], r["local_id"], r["name"], r["father_name"],
                               r["mobile"], r["address"], "By hand" if r["by_hand"] else "Post", r["status"],
                               dmy(r["expiry"]), r["total_paid"]) for r in rows], "jamaat_wise_subscribers")
    groups = {}
    for r in rows:
        groups.setdefault(r["jamaat"] or "(No Jamaat)", []).append(r)
    return render_template("report_jamaat.html", groups=groups, jam=jam, status=status, sort=sort, dir=d,
                           jamaats=jamaat_list(), total=len(rows))


@app.route("/reports/postal")
@login_required
def report_postal():
    """Address labels for posting the magazine: name, Jama'at, postal address; 12-16 per A4 page."""
    jam = jam_filter_value()
    status = request.args.get("status", "current")
    delivery = request.args.get("delivery", "post")
    try:
        per = int(request.args.get("per", 14))
    except ValueError:
        per = 14
    per = per if per in (12, 14, 16) else 14
    where, args = [], []
    scope(where, args)
    if jam == "none":
        where.append("jamaat_id IS NULL")
    elif jam and is_admin():
        where.append("jamaat_id=?")
        args.append(jam)
    if status == "current":
        where.append("status='Valid'")
    elif status:
        where.append("status=?")
        args.append(status)
    if delivery == "post":
        where.append("by_hand=0")
    elif delivery == "hand":
        where.append("by_hand=1")
    w = " WHERE " + " AND ".join(where) if where else ""
    rows = db.get_db().execute(
        f"SELECT * FROM v_subscribers{w} ORDER BY (jamaat IS NULL), lower(jamaat), lower(name), id",
        args).fetchall()
    if request.args.get("format") == "xlsx":
        return xlsx_response("Postal addresses", ["Name", "Jamaat", "Postal address"],
                             [(r["name"], r["jamaat"] or "", r["address"] or "") for r in rows], "postal_addresses")
    pages = [rows[i:i + per] for i in range(0, len(rows), per)]
    missing = sum(1 for r in rows if not (r["address"] or "").strip())
    return render_template("report_postal.html", pages=pages, per=per, rows_per_col=per // 2, jam=jam,
                           status=status, delivery=delivery, jamaats=jamaat_list(), total=len(rows), missing=missing)


@app.route("/reports/expiry")
@login_required
def report_expiry():
    """Expiry list with the amount each subscriber must pay to be active until the target 30 June."""
    mode = request.args.get("mode", "upcoming")
    t = today()
    d_from = request.args.get("from") or t.isoformat()
    d_to = request.args.get("to") or add_months(t, 2).isoformat()
    jam = jam_filter_value()
    extra = 1 if request.args.get("pay_to") == "next" else 0
    pay_to = "next" if extra else "fy"
    sort, d = get_sort("expiry")
    where, args = [], []
    scope(where, args)
    if mode == "expired":
        where.append("status='Expired'")
    else:
        where.append("status='Valid' AND expiry BETWEEN ? AND ?")
        args += [d_from, d_to]
    if jam and is_admin():
        where.append("jamaat_id=?")
        args.append(jam)
    found = db.get_db().execute(
        f"SELECT * FROM v_subscribers WHERE {' AND '.join(where)}{order_sql(sort, d)}", args).fetchall()
    rows = []
    for r in found:
        x = dict(r)
        q = renewal_quote(r["expiry"], extra)
        x["amt"], x["issues"], x["until"] = q["amount"], q["issues"], q["until"]
        rows.append(x)
    total_amt = sum(r["amt"] for r in rows)
    total_issues = sum(r["issues"] for r in rows)
    if request.args.get("format") == "xlsx":
        return xlsx_response("Expiry", ["ID", "Local ID", "Name", "Jamaat", "Mobile", "Address", "Expiry date",
                             "Issues to pay", "Amount to pay", "Active until"],
                             [(r["id"], r["local_id"], r["name"], r["jamaat"], r["mobile"], r["address"],
                               dmy(r["expiry"]), r["issues"], r["amt"], dmy(r["until"])) for r in rows],
                             "expiry_report")
    return render_template("report_expiry.html", rows=rows, mode=mode, d_from=d_from, d_to=d_to, jam=jam,
                           jamaats=jamaat_list(), sort=sort, dir=d, pay_to=pay_to,
                           total_amt=total_amt, total_issues=total_issues)


@app.route("/reports/yearly")
@login_required
def report_yearly():
    """Financial-year (July-June) collection report."""
    con = db.get_db()
    y = int(request.args.get("fy") or fy_of(today()))
    a, b = fy_range(y)
    sc, sa = "", []
    if not is_admin():
        sc, sa = " AND s.jamaat_id=?", [g.user["jamaat_id"]]
    monthly = {r["m"]: r for r in con.execute(
        f"""SELECT substr(p.paid_date, 1, 7) m, COUNT(*) n, SUM(p.amount) total
            FROM payments p JOIN subscribers s ON s.id=p.subscriber_id
            WHERE p.paid_date BETWEEN ? AND ?{sc} GROUP BY m""", [a, b] + sa)}
    months = []
    for i in range(12):
        dd = add_months(date(y, 7, 1), i)
        r = monthly.get(dd.strftime("%Y-%m"))
        months.append((dd.strftime("%B %Y"), r["n"] if r else 0, r["total"] if r else 0))
    by_jamaat = con.execute(
        f"""SELECT COALESCE(j.name,'(No Jamaat)') jamaat, COUNT(*) n, COUNT(DISTINCT p.subscriber_id) subs,
                   SUM(p.amount) total
            FROM payments p JOIN subscribers s ON s.id=p.subscriber_id LEFT JOIN jamaats j ON j.id=s.jamaat_id
            WHERE p.paid_date BETWEEN ? AND ?{sc} GROUP BY j.name ORDER BY j.name""", [a, b] + sa).fetchall()
    payments = con.execute(
        f"""SELECT p.*, s.name, s.local_id, COALESCE(j.name,'') jamaat FROM payments p
            JOIN subscribers s ON s.id=p.subscriber_id LEFT JOIN jamaats j ON j.id=s.jamaat_id
            WHERE p.paid_date BETWEEN ? AND ?{sc} ORDER BY p.paid_date, p.id""", [a, b] + sa).fetchall()
    if request.args.get("format") == "xlsx":
        return xlsx_response(f"FY {fy_label(y)}", ["Date", "Receipt no", "Book no", "ID", "Local ID", "Name",
                             "Jamaat", "Amount", "Expiry date"],
                             [(dmy(p["paid_date"]), p["receipt_no"], p["book_no"], p["subscriber_id"],
                               p["local_id"], p["name"], p["jamaat"], p["amount"], dmy(p["expiry_date"]))
                              for p in payments], f"payments_FY_{fy_label(y)}")
    all_years = {}
    for r in con.execute(f"""SELECT p.paid_date, p.amount, p.subscriber_id FROM payments p
                             JOIN subscribers s ON s.id=p.subscriber_id WHERE 1=1{sc}""", sa):
        e = all_years.setdefault(fy_of(r["paid_date"]), [0, 0.0, set()])
        e[0] += 1
        e[1] += r["amount"]
        e[2].add(r["subscriber_id"])
    summary = [(yy, v[0], len(v[2]), v[1]) for yy, v in sorted(all_years.items(), reverse=True)]
    return render_template("report_yearly.html", y=y, months=months, by_jamaat=by_jamaat,
                           payments=payments, total=sum(p["amount"] for p in payments),
                           summary=summary, years=[s[0] for s in summary] or [y])


# ------------------------------------------------------------------ Excel import (admin only)
# The uploaded sheet must have EVERY column of the template (the same fields as the web forms).
# Column order does not matter; capital letters, spaces and punctuation in the headings do not matter.
# Missing columns -> the file is rejected. Extra columns are ignored (and shown on the check page).
IMPORT_MAX_ROWS = 5000

IMPORT_SPECS = {
    "subscribers": {
        "label": "subscribers", "sheet": "Subscribers", "file": "subscribers_import_template",
        "cols": [("id", "ID"), ("local_id", "Local ID"), ("name", "Name"), ("father_name", "Father's Name"),
                 ("jamaat", "Jama'at"), ("address", "Address"), ("mobile", "Mobile"), ("email", "E-mail"),
                 ("by_hand", "Delivery"), ("active", "Active"), ("notes", "Notes")],
        "show": [("id", "ID"), ("local_id", "Local ID"), ("name", "Name"), ("jamaat", "Jama'at")],
        "info": [
            ("ID", "Optional", "Central subscriber ID. Leave empty to get the next free number. An ID already in the system is blocked.", "1301"),
            ("Local ID", "Optional", "Member number inside the Jama'at (must not repeat in one Jama'at).", "17"),
            ("Name", "REQUIRED", "Full name of the subscriber.", "Abdul Karim"),
            ("Father's Name", "Optional", "", "Md. Rahim"),
            ("Jama'at", "REQUIRED", "Jama'at name, spelled as in the system.", "Dhaka"),
            ("Address", "Optional", "Full postal address.", "House 5, Road 3, Mirpur, Dhaka"),
            ("Mobile", "Optional", "Keep the column as Text so the leading 0 stays.", "01711123456"),
            ("E-mail", "Optional", "", "karim@example.com"),
            ("Delivery", "Optional", "Post or By hand (empty = Post).", "Post"),
            ("Active", "Optional", "Yes or No (empty = Yes).", "Yes"),
            ("Notes", "Optional", "", ""),
        ],
        "example": [[1301, "17", "Abdul Karim", "Md. Rahim", "Dhaka", "House 5, Road 3, Mirpur, Dhaka",
                     "01711123456", "karim@example.com", "Post", "Yes", ""],
                    ["", "", "Rahima Begum", "Abdul Hamid", "Comilla", "Comilla Sadar", "01811123456", "",
                     "By hand", "Yes", ""]],
        "rules": [
            "The file must contain all the columns of the template (any order). A missing column rejects the whole file.",
            "Name and Jama'at are required in every row; Delivery must be Post or By hand; Active must be Yes or No.",
            "BLOCKED as double entry: an ID that is already used (when ID, name and Jama'at are all the same it is reported as a duplicate); "
            "the same Local ID twice in one Jama'at; the same name + Jama'at + mobile (or father's name); "
            "and any row that repeats an earlier row of the same file.",
            "A Jama'at name that does not exist yet is created when you confirm - the check page warns you first so you can fix spelling mistakes.",
        ],
    },
    "payments": {
        "label": "payments", "sheet": "Payments", "file": "payments_import_template",
        "cols": [("id", "ID"), ("receipt_no", "Receipt No"), ("book_no", "Book No"),
                 ("paid_date", "Payment Date"), ("amount", "Amount"), ("expiry_date", "Valid Until"),
                 ("note", "Note")],
        "show": [("id", "ID"), ("name", "Subscriber"), ("receipt_no", "Receipt No"), ("book_no", "Book No"),
                 ("paid_date", "Payment Date"), ("amount", "Amount"), ("expiry_date", "Valid Until")],
        "info": [
            ("ID", "REQUIRED", "Central subscriber ID. The subscriber must already exist.", "1301"),
            ("Receipt No", "REQUIRED", "Receipt number of this payment.", "15601"),
            ("Book No", "Optional", "Receipt book number.", "1000"),
            ("Payment Date", "REQUIRED", "Date of payment: an Excel date or DD-MM-YYYY.", "12-10-2026"),
            ("Amount", "REQUIRED", "Taka received (250 = 24 issues).", "250"),
            ("Valid Until", "REQUIRED", "Subscription valid until (Excel date or DD-MM-YYYY), on or after the payment date.", "30-09-2027"),
            ("Note", "Optional", "", "Cash"),
        ],
        "example": [[1301, "15601", "1000", "12-10-2026", 250, "30-09-2027", "Cash"]],
        "rules": [
            "The file must contain all the columns of the template (any order). A missing column rejects the whole file.",
            "ID, Receipt No, Payment Date, Amount and Valid Until are required in every row.",
            "BLOCKED as double entry: a row whose ID + Receipt No + Book No is already recorded, or that repeats an earlier row of the same file.",
            "You get a warning (not a block) if the same Receipt No + Book No was already used for a different subscriber.",
        ],
    },
}


class ImportProblem(Exception):
    """A problem with the uploaded file as a whole (shown to the user)."""


def _compact(s):
    return re.sub(r"[^a-z0-9]", "", str(s or "").lower())


def _cell_text(v):
    if v is None:
        return ""
    if isinstance(v, float) and v.is_integer():
        v = int(v)
    if isinstance(v, (datetime, date)):
        return v.strftime("%Y-%m-%d")
    return " ".join(str(v).replace("\r", " ").replace("\n", ", ").split())


def _parse_any_date(s):
    s = (s or "").strip()[:10]
    for fmt in ("%Y-%m-%d", "%d-%m-%Y", "%d/%m/%Y", "%d.%m.%Y"):
        try:
            return datetime.strptime(s, fmt).date()
        except ValueError:
            pass
    return None


def _parse_id(s):
    s = (s or "").strip()
    if not s:
        return None, None
    if re.fullmatch(r"\d+", s) and int(s) > 0:
        return int(s), None
    return None, "ID must be a whole number"


def _digits(s):
    return re.sub(r"\D", "", s or "")


def _clean_mobile(m):
    m = m.replace(" ", "").replace("-", "")
    if m.isdigit() and len(m) == 10 and m.startswith("1"):   # Excel dropped the leading 0 of 01XXXXXXXXX
        m = "0" + m
    return m


def parse_import_file(stream, kind):
    """Read the first sheet of an .xlsx and check it has every template column.
    Returns (rows, extra_column_headings)."""
    from openpyxl import load_workbook
    spec = IMPORT_SPECS[kind]
    wanted = {_compact(h): k for k, h in spec["cols"]}
    wb = load_workbook(stream, read_only=True, data_only=True)
    try:
        ws = wb.worksheets[0]
        header_map, extra, rows = None, [], []
        for i, row in enumerate(ws.iter_rows(values_only=True), start=1):
            if header_map is None:
                if not any(c not in (None, "") for c in row):
                    if i > 10:
                        break
                    continue                                  # skip empty lines above the headings
                m = {}
                for j, cell in enumerate(row):
                    k = wanted.get(_compact(cell))
                    if k and k not in m.values():
                        m[j] = k
                missing = [h for k, h in spec["cols"] if k not in m.values()]
                if missing:
                    raise ImportProblem(
                        f"This file does not match the {spec['label']} template. Missing column(s): "
                        f"{', '.join(missing)}. The first row of the sheet must have the headings: "
                        f"{', '.join(h for _k, h in spec['cols'])}. Download the template and copy your data into it.")
                header_map = m
                extra = [str(c).strip() for j, c in enumerate(row) if c not in (None, "") and j not in m]
                continue
            rec = {k: "" for k, _h in spec["cols"]}
            for j, k in header_map.items():
                if j < len(row):
                    rec[k] = _cell_text(row[j])
            if not any(rec.values()):
                continue
            if len(rows) >= IMPORT_MAX_ROWS:
                raise ImportProblem(f"The file has more than {IMPORT_MAX_ROWS} rows. Please upload it in smaller parts.")
            rec["_row"] = i
            rows.append(rec)
    finally:
        wb.close()
    if header_map is None:
        raise ImportProblem("The first sheet of this file is empty.")
    return rows, extra


def _summarize(results, spec, **extra):
    ok = [x for x in results if x["status"] == "ok"]
    out = {"results": results, "ok": len(ok), "skipped": len(results) - len(ok),
           "warnings": sum(1 for x in ok if x["warn"]), "total": len(results), "columns": spec["show"]}
    out.update(extra)
    return out


def process_subscribers(rows, commit=False):
    """Check every row (and save the good ones when commit=True)."""
    spec = IMPORT_SPECS["subscribers"]
    con = db.get_db()
    jam_ids = {r["name"].lower(): r["id"] for r in con.execute("SELECT id, name FROM jamaats")}
    seen_id, seen_local, seen_person, new_jamaats, results = set(), set(), set(), {}, []
    for r in rows:
        res = {"row": r["_row"], "id": r["id"], "local_id": r["local_id"], "name": r["name"],
               "jamaat": r["jamaat"], "status": "ok", "notes": [], "warn": False}
        results.append(res)

        def block(msg):
            res["status"] = "skip"
            res["notes"].append(msg)

        def warn(msg):
            res["warn"] = True
            res["notes"].append(msg)

        name, jname, local_id = r["name"], r["jamaat"].strip(), r["local_id"]
        father, mobile, email = r["father_name"], _clean_mobile(r["mobile"]), r["email"]
        errs = []
        if not name:
            errs.append("Name is missing")
        if not jname:
            errs.append("Jama'at is missing")
        sid, e = _parse_id(r["id"])
        if e:
            errs.append(e)
        d = r["by_hand"].strip().lower()
        by_hand = 1 if d in ("by hand", "byhand", "hand") else 0
        if d not in ("", "post", "by hand", "byhand", "hand"):
            errs.append("Delivery must be Post or By hand")
        a = r["active"].strip().lower()
        active = 0 if a in ("no", "n", "false", "0") else 1
        if a not in ("", "yes", "y", "true", "1", "no", "n", "false", "0"):
            errs.append("Active must be Yes or No")
        if errs:
            res["status"], res["notes"] = "skip", errs
            continue
        # ---- the central ID
        if sid is not None:
            if sid in seen_id:
                block(f"ID {sid} appears twice in this file")
                continue
            seen_id.add(sid)
            ex = con.execute("SELECT s.id, s.name, j.name AS jamaat FROM subscribers s "
                             "LEFT JOIN jamaats j ON j.id = s.jamaat_id WHERE s.id=?", (sid,)).fetchone()
            if ex:
                ex_j = ex["jamaat"] or "no Jama'at"
                if ex["name"].lower() == name.lower() and (ex["jamaat"] or "").lower() == jname.lower():
                    block(f"Duplicate - ID {sid}, name and Jama'at are already in the system")
                else:
                    block(f"ID {sid} is already used by {ex['name']} ({ex_j})")
                continue
        # ---- which Jama'at
        jid = jam_ids.get(jname.lower())
        if jid is None:
            e2 = new_jamaats.setdefault(jname.lower(), [jname, 0])
            e2[1] += 1
            warn("New Jama'at will be created")
        jkey = jid if jid is not None else "new:" + jname.lower()
        # ---- Local ID must be unique inside a Jama'at
        if local_id:
            lk = (jkey, local_id.lower())
            if lk in seen_local:
                block(f"Local ID {local_id} is repeated in this file")
                continue
            if jid is not None:
                c = con.execute("SELECT id, name FROM subscribers WHERE jamaat_id=? AND lower(local_id)=lower(?)",
                                (jid, local_id)).fetchone()
                if c:
                    block(f"Local ID {local_id} already belongs to #{c['id']} {c['name']}")
                    continue
            seen_local.add(lk)
        # ---- the same person already in the system, or twice in the file
        pk = (jkey, name.lower(), _digits(mobile) or father.lower(), local_id.lower())
        if pk in seen_person:
            block("The same member appears twice in this file")
            continue
        seen_person.add(pk)
        if jid is not None:
            cands = con.execute("SELECT id, local_id, father_name, mobile FROM subscribers "
                                "WHERE jamaat_id=? AND lower(name)=lower(?)", (jid, name)).fetchall()
            dup = None
            for c in cands:
                if mobile and c["mobile"] and _digits(mobile) == _digits(c["mobile"]):
                    dup = c
                elif father and c["father_name"] and father.lower() == c["father_name"].lower():
                    dup = c
            if dup:
                block(f"Duplicate - already in the system as #{dup['id']} (same name, Jama'at and mobile / father's name)")
                continue
            if cands and not mobile and not father:
                same = [c for c in cands if not c["local_id"] or not local_id]
                if same:
                    block(f"A member with the same name and Jama'at already exists (#{same[0]['id']}) and this row has no "
                          f"mobile or father's name to tell them apart - add by hand if it is a different person")
                    continue
            if cands:
                warn(f"Another member with the same name exists (#{cands[0]['id']}) - added anyway")
        if email and "@" not in email:
            warn("E-mail looks wrong - left blank")
            email = ""
        if commit:
            vals = [local_id or None, name, father or None, get_or_create_jamaat(jname), r["address"] or None,
                    mobile or None, email or None, by_hand, active, r["notes"] or None]
            cols = "local_id,name,father_name,jamaat_id,address,mobile,email,by_hand,active,notes"
            if sid is not None:                      # keep the ID given in the file; otherwise the next free number
                cols, vals = "id," + cols, [sid] + vals
            con.execute(f"INSERT INTO subscribers({cols}) VALUES ({','.join('?' * len(vals))})", vals)
    return _summarize(results, spec, new_jamaats=sorted(new_jamaats.values()))


def process_payments(rows, commit=False):
    """Check every payment row (and save the good ones when commit=True)."""
    spec = IMPORT_SPECS["payments"]
    con = db.get_db()
    names = {r["id"]: r["name"] for r in con.execute("SELECT id, name FROM subscribers")}
    keys, owner = set(), {}
    for r in con.execute("SELECT subscriber_id, COALESCE(receipt_no,'') rc, COALESCE(book_no,'') bk FROM payments"):
        keys.add((r["subscriber_id"], r["rc"].lower(), r["bk"].lower()))
        owner.setdefault((r["rc"].lower(), r["bk"].lower()), r["subscriber_id"])
    seen, results, total = set(), [], 0.0
    t = today()
    for r in rows:
        res = {"row": r["_row"], "id": r["id"], "name": "", "receipt_no": r["receipt_no"], "book_no": r["book_no"],
               "paid_date": r["paid_date"], "amount": r["amount"], "expiry_date": r["expiry_date"],
               "status": "ok", "notes": [], "warn": False}
        results.append(res)
        errs = []
        sid, e = _parse_id(r["id"])
        if e:
            errs.append(e)
        elif sid is None:
            errs.append("ID is missing")
        elif sid not in names:
            errs.append(f"No subscriber with ID {sid}")
        else:
            res["name"] = names[sid]
        receipt, book = r["receipt_no"].strip(), r["book_no"].strip()
        if not receipt:
            errs.append("Receipt No is missing")
        paid, exp = _parse_any_date(r["paid_date"]), _parse_any_date(r["expiry_date"])
        if not paid:
            errs.append("Payment Date is missing or not a date (use DD-MM-YYYY)")
        else:
            res["paid_date"] = paid.strftime("%d-%m-%Y")
        if not exp:
            errs.append("Valid Until is missing or not a date (use DD-MM-YYYY)")
        else:
            res["expiry_date"] = exp.strftime("%d-%m-%Y")
        try:
            amount = float(r["amount"].replace(",", ""))
            if amount <= 0:
                raise ValueError
        except ValueError:
            amount = None
            errs.append("Amount must be a number above 0")
        if paid and exp and exp < paid:
            errs.append("Valid Until is before the Payment Date")
        if not errs:
            key = (sid, receipt.lower(), book.lower())
            if key in keys:
                errs.append("Duplicate - this ID, Receipt No and Book No are already recorded")
            elif key in seen:
                errs.append("Duplicate - this ID, Receipt No and Book No appear twice in this file")
        if errs:
            res["status"], res["notes"] = "skip", errs
            continue
        seen.add(key)
        o = owner.get((receipt.lower(), book.lower()))
        if o is not None and o != sid:
            res["warn"] = True
            res["notes"].append(f"Same Receipt No and Book No was already used for subscriber #{o}")
        if exp > add_months(paid, 60):
            res["warn"] = True
            res["notes"].append("Valid Until is more than 5 years after the payment - please check")
        if paid > t:
            res["warn"] = True
            res["notes"].append("Payment Date is in the future")
        total += amount
        if commit:
            con.execute("""INSERT INTO payments(subscriber_id,receipt_no,book_no,paid_date,amount,expiry_date,note)
                           VALUES (?,?,?,?,?,?,?)""",
                        (sid, receipt, book or None, paid.isoformat(), amount, exp.isoformat(), r["note"] or None))
    return _summarize(results, spec, amount=total)


PROCESSORS = {"subscribers": process_subscribers, "payments": process_payments}


def _pending_dir():
    base = tempfile.gettempdir() if db.IS_PG else os.path.dirname(os.path.abspath(db.DB_PATH))
    d = os.path.join(base, "tmp_imports")
    os.makedirs(d, exist_ok=True)
    return d


def save_pending(kind, rows):
    d, now = _pending_dir(), time.time()
    for fn in os.listdir(d):                       # tidy: forget uploads older than one hour
        p = os.path.join(d, fn)
        try:
            if now - os.path.getmtime(p) > 3600:
                os.remove(p)
        except OSError:
            pass
    token = secrets.token_hex(16)
    with open(os.path.join(d, token + ".json"), "w", encoding="utf-8") as fh:
        json.dump({"uid": g.user["id"], "kind": kind, "rows": rows}, fh)
    return token


def load_pending(token, kind):
    if not re.fullmatch(r"[0-9a-f]{32}", token or "") or session.get("import_token") != token:
        return None
    try:
        with open(os.path.join(_pending_dir(), token + ".json"), encoding="utf-8") as fh:
            data = json.load(fh)
    except (OSError, ValueError):
        return None
    return data["rows"] if data.get("uid") == g.user["id"] and data.get("kind") == kind else None


def drop_pending(token):
    try:
        os.remove(os.path.join(_pending_dir(), token + ".json"))
    except OSError:
        pass
    session.pop("import_token", None)


def _spec(kind):
    if kind not in IMPORT_SPECS:
        abort(404)
    return IMPORT_SPECS[kind]


@app.errorhandler(413)
def too_big(_e):
    flash("That file is too big (limit 5 MB).", "danger")
    return redirect(url_for("import_page", kind="subscribers"))


@app.route("/import/<kind>", methods=["GET", "POST"])
@admin_required
def import_page(kind):
    spec = _spec(kind)
    if request.method == "POST":
        f = request.files.get("file")
        if not f or not f.filename:
            flash("Choose an Excel file first.", "danger")
            return redirect(url_for("import_page", kind=kind))
        if not f.filename.lower().endswith((".xlsx", ".xlsm")):
            flash("Please upload an .xlsx file (in Excel: File > Save As > Excel Workbook).", "danger")
            return redirect(url_for("import_page", kind=kind))
        try:
            rows, extra = parse_import_file(f.stream, kind)
        except ImportProblem as e:
            flash(str(e), "danger")
            return redirect(url_for("import_page", kind=kind))
        except Exception:
            flash("This file could not be read. Please upload a normal .xlsx file (not password protected).", "danger")
            return redirect(url_for("import_page", kind=kind))
        if not rows:
            flash("The file has a heading row but no data rows under it.", "danger")
            return redirect(url_for("import_page", kind=kind))
        token = save_pending(kind, rows)
        session["import_token"] = token
        return render_template("import_preview.html", kind=kind, spec=spec, token=token, extra=extra,
                               filename=f.filename, result=PROCESSORS[kind](rows))
    return render_template("import.html", kind=kind, spec=spec)


@app.route("/import/<kind>/confirm", methods=["POST"])
@admin_required
def import_confirm(kind):
    spec = _spec(kind)
    token = request.form.get("token", "")
    rows = load_pending(token, kind)
    if rows is None:
        flash("This upload has expired. Please upload the file again.", "warning")
        return redirect(url_for("import_page", kind=kind))
    result = PROCESSORS[kind](rows, commit=True)
    db.sync_sequences(db.get_db())
    db.get_db().commit()
    drop_pending(token)
    return render_template("import_done.html", kind=kind, spec=spec, result=result)


@app.route("/import/<kind>/template")
@admin_required
def import_template(kind):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.worksheet.datavalidation import DataValidation
    spec = _spec(kind)
    heads = [h for _k, h in spec["cols"]]
    green = PatternFill("solid", fgColor="1F6F43")
    wb = Workbook()
    ws = wb.active                                   # the FIRST sheet is the one that is read on upload
    ws.title = spec["sheet"]
    ws.append(heads)
    widths = {"ID": 8, "Local ID": 11, "Name": 28, "Father's Name": 26, "Jama'at": 18, "Address": 40, "Mobile": 15,
              "E-mail": 26, "Delivery": 11, "Active": 9, "Notes": 26, "Receipt No": 13, "Book No": 11,
              "Payment Date": 15, "Amount": 10, "Valid Until": 14, "Note": 24}
    for i, h in enumerate(heads, start=1):
        ws.cell(row=1, column=i).font = Font(bold=True, color="FFFFFF")
        ws.cell(row=1, column=i).fill = green
        letter = ws.cell(row=1, column=i).column_letter
        ws.column_dimensions[letter].width = widths.get(h, 16)
        if h in ("Local ID", "Mobile", "Receipt No", "Book No"):
            ws.column_dimensions[letter].number_format = "@"          # text: keeps leading zeros
        if h in ("Payment Date", "Valid Until"):
            ws.column_dimensions[letter].number_format = "DD-MM-YYYY"
    ws.freeze_panes = "A2"
    if kind == "subscribers":
        dv1 = DataValidation(type="list", formula1='"Post,By hand"', allow_blank=True)
        dv2 = DataValidation(type="list", formula1='"Yes,No"', allow_blank=True)
        ws.add_data_validation(dv1)
        ws.add_data_validation(dv2)
        dv1.add("I2:I2000")
        dv2.add("J2:J2000")
    ex = wb.create_sheet("Example")                  # filled-in sample, NOT uploaded
    ex.append(heads)
    for row in spec["example"]:
        ex.append(row)
    for i, h in enumerate(heads, start=1):
        ex.cell(row=1, column=i).font = Font(bold=True, color="FFFFFF")
        ex.cell(row=1, column=i).fill = PatternFill("solid", fgColor="66727F")
        ex.column_dimensions[ex.cell(row=1, column=i).column_letter].width = widths.get(h, 16)
    guide = wb.create_sheet("How to fill")
    guide.append(["Column", "Required?", "What to write", "Example"])
    for row in spec["info"]:
        guide.append(list(row))
    guide.append([])
    guide.append(["Rules"])
    for rule in spec["rules"]:
        guide.append(["", "", rule])
    guide.append(["", "", "Type your data on the first sheet, from row 2. Do not change the headings in row 1. "
                          "The 'Example' and 'How to fill' sheets are only for your information and are not uploaded."])
    for c in guide[1]:
        c.font = Font(bold=True)
    for col, w in zip("ABCD", (16, 12, 100, 34)):
        guide.column_dimensions[col].width = w
    for row in guide.iter_rows():
        for c in row:
            c.alignment = Alignment(wrap_text=True, vertical="top")
    buf = io.BytesIO()
    wb.save(buf)
    return Response(buf.getvalue(), headers={
        "Content-Disposition": f"attachment; filename={spec['file']}.xlsx",
        "Content-Type": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"})


# ------------------------------------------------------------------ command line
@app.cli.command("init-db")
def cli_init():
    """Create / upgrade the database tables."""
    db.init_db()
    click.echo("Database ready.")


@app.cli.command("create-user")
@click.argument("username")
@click.option("--role", type=click.Choice(["admin", "jamaat"]), default="admin", show_default=True)
@click.option("--jamaat", "jamaat_name", default=None, help="Jama'at name (required for role jamaat)")
@click.password_option()
def cli_user(username, role, jamaat_name, password):
    """Create or reset a login user."""
    db.init_db()
    con = db.connect()
    jid = None
    if role == "jamaat":
        r = con.execute("SELECT id FROM jamaats WHERE lower(name)=lower(?)", (jamaat_name or "",)).fetchone()
        if not r:
            raise click.ClickException("Give an existing Jama'at name with --jamaat")
        jid = r["id"]
    con.execute("""INSERT INTO users(username,password_hash,role,jamaat_id) VALUES (?,?,?,?)
                   ON CONFLICT(username) DO UPDATE SET password_hash=excluded.password_hash,
                   role=excluded.role, jamaat_id=excluded.jamaat_id""",
                (username, generate_password_hash(password), role, jid))
    con.commit()
    click.echo(f"User '{username}' saved.")


@app.cli.command("import-access")
@click.argument("mdb_file", type=click.Path(exists=True))
@click.option("--force", is_flag=True, help="Wipe existing data first.")
def cli_import(mdb_file, force):
    """Import subscribers and payments from the old Access .mdb file."""
    import import_access
    click.echo(import_access.run(mdb_file, force=force))


@app.route("/healthz")
def healthz():
    return "ok"


def ensure_admin():
    """Online start-up helper: if there is no user yet and ADMIN_USERNAME / ADMIN_PASSWORD are set, create that admin."""
    user, pw = os.environ.get("ADMIN_USERNAME"), os.environ.get("ADMIN_PASSWORD")
    con = db.connect()
    if user and pw and con.execute("SELECT COUNT(*) AS n FROM users").fetchone()["n"] == 0:
        con.execute("INSERT INTO users(username,password_hash,role) VALUES (?,?,'admin')",
                    (user, generate_password_hash(pw)))
        con.commit()
        print(f"First admin '{user}' created.")
    con.close()


@app.cli.command("migrate-sqlite")
@click.argument("sqlite_file", type=click.Path(exists=True))
@click.option("--force", is_flag=True, help="Wipe the PostgreSQL data first.")
def cli_migrate(sqlite_file, force):
    """Copy everything (subscribers, payments, users) from the old SQLite file into PostgreSQL."""
    if not db.IS_PG:
        raise click.ClickException("Set DATABASE_URL to the PostgreSQL address first.")
    db.init_db()
    src = db.connect(sqlite_file)
    dst = db.connect()
    if dst.execute("SELECT COUNT(*) AS n FROM subscribers").fetchone()["n"] and not force:
        raise click.ClickException("PostgreSQL already has subscribers. Use --force to wipe and copy again.")
    if force:
        for t in ("payments", "subscribers", "users", "jamaats"):
            dst.execute(f"DELETE FROM {t}")
    plan = [("jamaats", "id,name"),
            ("users", "id,username,password_hash,role,jamaat_id"),
            ("subscribers", "id,local_id,name,father_name,jamaat_id,address,mobile,email,by_hand,active,notes,created_at"),
            ("payments", "id,subscriber_id,receipt_no,book_no,paid_date,amount,expiry_date,note,created_at")]
    for table, cols in plan:
        n = 0
        for r in src.execute(f"SELECT {cols} FROM {table} ORDER BY id"):
            dst.execute(f"INSERT INTO {table}({cols}) VALUES ({','.join('?' * len(r))})", tuple(r))
            n += 1
        click.echo(f"{table}: {n} rows copied")
    db.sync_sequences(dst)
    dst.commit()
    click.echo("Done.")


if __name__ == "__main__":
    db.init_db()
    app.run(debug=True)
